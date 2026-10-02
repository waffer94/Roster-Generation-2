#!/usr/bin/env python3
"""
build_signoff.py
================
Generate TLMS sign-off sheets (.xlsx) from an Excel export + a keyword mapping.

Usage:
    python build_signoff.py EXPORT.xlsx MAPPING.xlsx [--out OUTPUT_DIR]

For every data row in the export it:
  1. Matches the course name against the keyword mapping to find the TLMS course name(s).
  2. Parses participants (handling First\\tLast\\temail, name-only, bulleted formats).
  3. Creates a .xlsx with Sheet1 (user import) and Sheet2 (course completion).
  4. Saves as "TLMS Sign Off - YYYYMMDD - Company - CODE(s).xlsx".
"""

import argparse
import os
import re
import sys
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment

# --------------------------------------------------------------------------- #
# Mapping loader
# --------------------------------------------------------------------------- #

def load_mapping(mapping_path: str) -> list:
    """Load the keyword -> TLMS course mapping from the 'List - Sign off sheets'
    sheet.  Returns a list of dicts sorted longest-keyword-first."""
    wb = load_workbook(mapping_path, data_only=True)
    # find the right sheet
    target = None
    for sn in wb.sheetnames:
        if "sign off" in sn.lower():
            target = sn
            break
    if not target:
        sys.exit(f"No 'Sign off sheets' tab found in {mapping_path}")
    ws = wb[target]
    rows = []
    for r in range(2, ws.max_row + 1):
        kw = (ws.cell(row=r, column=2).value or "").strip()
        all_p = (ws.cell(row=r, column=3).value or "").strip()
        part_p = (ws.cell(row=r, column=4).value or "").strip()
        recert_p = (ws.cell(row=r, column=5).value or "").strip()
        if not kw:
            continue
        rows.append({
            "keyword": kw,
            "all": all_p,          # single course for everyone
            "part_course": part_p, # course for Participant Names col
            "recert_course": recert_p,  # course for Recertification col
        })
    # sort by keyword length descending so longer/more-specific keywords match first
    rows.sort(key=lambda r: len(r["keyword"]), reverse=True)
    return rows


def match_mapping(course_name: str, mapping: list) -> dict | None:
    """Find the first mapping entry whose keyword appears in the course name."""
    low = course_name.lower()
    for entry in mapping:
        if entry["keyword"].lower() in low:
            return entry
    return None


def _extract_code(tlms_course: str) -> str:
    """'Blended Intermediate First Aid & CPR Level C (TIFA-B-2-ON)' -> 'TIFAB'."""
    m = re.search(r"\(([^)]+)\)", tlms_course)
    if not m:
        return ""
    raw = m.group(1)
    return raw.replace("-ON", "").replace("-2", "").replace("-", "")


# --------------------------------------------------------------------------- #
# Field extraction (reuses patterns from build_rosters)
# --------------------------------------------------------------------------- #

def _strip_acct(s: str) -> str:
    s = re.sub(r"^\s*\d+\s*", "", s)
    s = re.sub(r"\s*\d+\s*$", "", s)
    return re.sub(r"\s+", " ", s).strip(" :")


def clean_company(customer: str, name_col: str = "") -> str:
    a = (name_col or "").strip()
    if " : " in a:
        first_seg = a.split(" : ")[0].strip()
        parent = _strip_acct(first_seg)
        if parent:
            return parent
    c = (customer or "").strip()
    if ":" in c:
        segs = [s.strip() for s in c.split(":")]
        c = max(segs, key=lambda s: sum(ch.isalpha() for ch in s))
    return _strip_acct(c)


def extract_course_name(name_col: str) -> str:
    raw = (name_col or "").strip()
    m = re.search(r"\b(\d{8})\s+", raw)
    if m:
        after = raw[m.end():]
    else:
        after = raw.split(" : ", 1)[1] if " : " in raw else raw
        after = re.sub(r"^\s*\d{8}\s+", "", after).strip()
    after = re.split(r"_x000D_|[\r\n]", after)[0].strip()
    return re.sub(r"\s+", " ", after).strip()


def _as_dt(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, str) and v.strip():
        for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%Y%m%d"):
            try:
                return datetime.strptime(v.strip(), fmt)
            except ValueError:
                pass
    return None


# --------------------------------------------------------------------------- #
# Participant parsing — handles tab-separated (first\tlast\temail), plain
# name-per-line, bulleted, and multi-line-with-email formats
# --------------------------------------------------------------------------- #

def _parse_signoff_participants(blob: str) -> list:
    """Returns list of (firstname, lastname, email) tuples."""
    if not blob:
        return []
    text = str(blob).replace("_x000D_", "")
    results = []
    for line in re.split(r"[\r\n]+", text):
        line = line.strip()
        if not line or line.lower() == "tbd":
            continue
        line = re.sub(r"^[•·▪‣◦*\-]\s*", "", line)  # strip bullets
        parts = line.split("\t")
        if len(parts) >= 3:
            # first\tlast\temail
            first = parts[0].strip()
            last = parts[1].strip()
            email = parts[2].strip()
            if first or last:
                results.append((first, last, email))
        elif len(parts) == 2:
            # first\tlast  (no email)
            first = parts[0].strip()
            last = parts[1].strip()
            if first or last:
                results.append((first, last, ""))
        else:
            # space-separated: "First Last" or "First Last email@..."
            tokens = line.split()
            email_tok = ""
            name_toks = []
            for t in tokens:
                if "@" in t:
                    email_tok = t
                else:
                    name_toks.append(t)
            if name_toks:
                first = name_toks[0]
                last = " ".join(name_toks[1:]) if len(name_toks) > 1 else ""
                results.append((first, last, email_tok))
    return results


# --------------------------------------------------------------------------- #
# Output generation
# --------------------------------------------------------------------------- #

SHEET1_HEADERS = ["Login", "Firstname", "Lastname", "Email", "Company",
                  "Instructor (User)", "Branch", "Group", "Course"]
SHEET2_HEADERS = ["Usertocourses", "course", "EnrolledOndate",
                  "CompletionDate", "Status"]


def _make_login(email: str, first: str, last: str) -> str:
    """Use email as login; if no email, construct a placeholder."""
    if email:
        return email
    return f"{first.lower()}.{last.lower()}@placeholder.com" if first and last else ""


def build_signoff(rec, mapping, out_dir):
    """Build one sign-off .xlsx from a single export row."""
    name_col = rec.get("name") or ""
    company = clean_company(rec.get("customer"), name_col)
    course_name = extract_course_name(name_col)
    instructor = (rec.get("instructor") or "").strip()
    start_dt = _as_dt(rec.get("start"))
    end_dt = _as_dt(rec.get("end"))

    if not start_dt:
        return ("skipped", "no course date", None)

    # company field in output: "Company - YYYY-MM-DD"
    company_field = f"{company} - {start_dt.strftime('%Y-%m-%d')}" if company else \
                    start_dt.strftime('%Y-%m-%d')

    # enrolled/completion date = end date + 1 day
    enroll_dt = (end_dt or start_dt) + timedelta(days=1)
    enroll_str = enroll_dt.strftime("%Y-%m-%d 00:00:00")

    # match mapping
    entry = match_mapping(course_name, mapping)

    # parse participants
    part_people = _parse_signoff_participants(rec.get("participants"))
    recert_people = _parse_signoff_participants(rec.get("recert_participants"))

    # Build rows: each person gets (first, last, email, tlms_course)
    rows = []
    if entry:
        if entry["all"]:
            # single course for everyone
            tlms = entry["all"]
            for f, l, e in part_people:
                rows.append((f, l, e, tlms))
            for f, l, e in recert_people:
                rows.append((f, l, e, tlms))
        else:
            # separate courses
            for f, l, e in part_people:
                rows.append((f, l, e, entry["part_course"]))
            for f, l, e in recert_people:
                rows.append((f, l, e, entry["recert_course"]))
    else:
        # no mapping match — use raw course name
        for f, l, e in part_people:
            rows.append((f, l, e, course_name))
        for f, l, e in recert_people:
            rows.append((f, l, e, course_name))

    if not rows:
        return ("skipped", "no participants", None)

    # sort by lastname, firstname
    rows.sort(key=lambda r: (r[1].lower(), r[0].lower()))

    # Create workbook
    wb = Workbook()
    hdr_font = Font(name="Arial", bold=True, size=11)
    data_font = Font(name="Arial", size=11)

    # Sheet1
    ws1 = wb.active
    ws1.title = "Sheet1"
    for ci, h in enumerate(SHEET1_HEADERS, 1):
        c = ws1.cell(row=1, column=ci, value=h)
        c.font = hdr_font
    for ri, (first, last, email, tlms) in enumerate(rows, 2):
        login = _make_login(email, first, last)
        vals = [login, first, last, email, company_field, instructor,
                "fastrescue", "", tlms]
        for ci, v in enumerate(vals, 1):
            c = ws1.cell(row=ri, column=ci, value=v)
            c.font = data_font

    # Sheet2
    ws2 = wb.create_sheet("Sheet2")
    for ci, h in enumerate(SHEET2_HEADERS, 1):
        c = ws2.cell(row=1, column=ci, value=h)
        c.font = hdr_font
    for ri, (first, last, email, tlms) in enumerate(rows, 2):
        login = _make_login(email, first, last)
        vals = [login, tlms, enroll_str, enroll_str, "completed"]
        for ci, v in enumerate(vals, 1):
            c = ws2.cell(row=ri, column=ci, value=v)
            c.font = data_font

    # auto-width columns
    for ws in (ws1, ws2):
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    # Filename: TLMS Sign Off - YYYYMMDD - Company - CODE(s).xlsx
    date_str = start_dt.strftime("%Y%m%d")
    codes = set()
    for _, _, _, tlms in rows:
        code = _extract_code(tlms)
        if code:
            codes.add(code)
    code_str = " & ".join(sorted(codes)) if codes else ""
    parts = ["TLMS Sign Off", date_str, company]
    if code_str:
        parts.append(code_str)
    fname = re.sub(r'[\\/:*?"<>|]', "", " - ".join(p for p in parts if p)) + ".xlsx"

    out_path = os.path.join(out_dir, fname)
    wb.save(out_path)
    return ("ok", out_path, len(rows))


# --------------------------------------------------------------------------- #
# Export reader
# --------------------------------------------------------------------------- #

SH = {
    "name": "name", "customer": "customer",
    "participant names": "participants",
    "recertification participants": "recert_participants",
    "course date": "start", "course end date": "end",
    "instructors": "instructor",
}


def read_signoff_export(xlsx_path):
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    headers = {}
    for c in range(1, ws.max_column + 1):
        h = (ws.cell(row=1, column=c).value or "").strip().lower()
        if h in SH:
            headers[SH[h]] = c
    rows = []
    for r in range(2, ws.max_row + 1):
        if not any(ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)):
            continue
        rows.append({k: ws.cell(row=r, column=col).value for k, col in headers.items()})
    return rows


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #

def main():
    ap = argparse.ArgumentParser(description="Build TLMS sign-off sheets from an Excel export.")
    ap.add_argument("export", help="Path to the sign-off export (.xlsx)")
    ap.add_argument("mapping", help="Path to the mapping workbook (with 'List - Sign off sheets' tab)")
    ap.add_argument("--out", default="signoff_sheets", help="Output folder")
    args = ap.parse_args()

    os.makedirs(args.out, exist_ok=True)
    mapping = load_mapping(args.mapping)
    print(f"Loaded {len(mapping)} mapping entries\n")

    rows = read_signoff_export(args.export)
    print(f"{len(rows)} course(s) in export\n")
    made, skipped = 0, 0
    for rec in rows:
        status, info, n = build_signoff(rec, mapping, args.out)
        if status == "ok":
            made += 1
            print(f"  {n:2d} participants -> {os.path.basename(info)}")
        else:
            skipped += 1
            print(f"  [SKIPPED] {info}")
    print(f"\nDone: {made} sign-off sheet(s) created"
          + (f", {skipped} skipped" if skipped else ""))


if __name__ == "__main__":
    main()
